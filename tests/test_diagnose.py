import json

from scf_api.cli import main


def test_diagnose_clean_workbook(clean_workbook, capsys):
    assert main(["diagnose", "--workbook", str(clean_workbook), "--scf-version", "2099.1.1"]) == 0
    out = capsys.readouterr().out
    assert "PARSED SCF 2099.1.1" in out and "controls=6" in out


def test_diagnose_unknown_tab_names_the_closest_tab(workbook_factory, capsys):
    path = workbook_factory("unknown-tab")
    assert main(["diagnose", "--workbook", str(path), "--scf-version", "2099.1.1"]) == 1
    out = capsys.readouterr().out
    assert "no tab matches kind 'sources'" in out
    assert "SHEET_PATTERNS" in out and "Reference Register" in out


def test_diagnose_missing_column_suggests_headers(clean_workbook, tmp_path, capsys):
    from openpyxl import load_workbook

    wb = load_workbook(clean_workbook)
    ws = wb["Evidence Request List 2099.1"]
    ws["D1"] = "Proof Item"  # was "Documentation Artifact"
    wb.save(clean_workbook)
    assert main(["diagnose", "--workbook", str(clean_workbook), "--scf-version", "2099.1.1"]) == 1
    out = capsys.readouterr().out
    assert "no column matching" in out and "alts=(...)" in out


def test_triage_groups_findings(tmp_path, capsys):
    report = {
        "scf_version": "2099.1.1",
        "summary": {"fail": 3, "warn": 0, "resolved": 0, "ok": False},
        "findings": [
            {"check": "control-risk", "severity": "fail", "message": f"bad {i}", "owner": "X"}
            for i in range(3)
        ],
    }
    path = tmp_path / "r.json"
    path.write_text(json.dumps(report))
    assert main(["triage", "--report", str(path), "--limit", "1"]) == 1
    out = capsys.readouterr().out
    assert "control-risk: 3 (fail=3)" in out and "hint:" in out and out.count("bad ") == 1


def test_headers_lists_columns_with_samples(clean_workbook, capsys):
    assert (
        main(
            [
                "headers",
                "--workbook",
                str(clean_workbook),
                "--sheet",
                "SCF 2099",
                "--grep",
                "risk",
                "--sample",
            ]
        )
        == 0
    )
    assert "Risk Threat Summary" in capsys.readouterr().out
