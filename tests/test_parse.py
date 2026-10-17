from scf_api.parse import parse_workbook

VERSION = "2099.1.1"


def test_parse_structure(clean_workbook):
    result = parse_workbook(clean_workbook, VERSION)

    assert result.catalog_name == "scf-2099-1-1.json"
    assert set(result.files) == {
        "scf-2099-1-1.json",
        "scf-assessment-objectives.json",
        "scf-compensating-controls.json",
        "scf-crosswalks.json",
        "scf-evidence-requests.json",
        "scf-privacy-principles.json",
        "scf-risks.json",
        "scf-threats.json",
    }

    catalog = result.catalog
    assert catalog["framework"] == {"id": "scf-2099-1-1", "name": "SCF", "version": VERSION}
    assert list(catalog["families"]) == ["GOV", "MON", "TST"]
    assert catalog["families"]["GOV"] == "Governance"

    controls = catalog["controls"]
    assert [c["control_id"] for c in controls] == [
        "GOV-01",
        "GOV-01.1",
        "GOV-02",
        "MON-01",
        "MON-02",
        "TST-01",
    ]
    gov01 = controls[0]
    # Field order is part of the output contract.
    assert list(gov01) == [
        "control_id",
        "title",
        "family",
        "description",
        "scf_question",
        "relative_weight",
        "conformity_cadence",
        "evidence_requests",
        "pptdf",
        "nist_csf_function",
        "scrm_focus",
        "maturity",
        "profiles",
        "possible_solutions",
        "risks",
        "threats",
        "errata",
    ]
    assert gov01["family"] == "GOV"
    assert gov01["relative_weight"] == 10
    assert gov01["evidence_requests"] == ["E-GOV-01", "E-GOV-02"]
    assert gov01["scrm_focus"] == {"strategic": True, "operational": True, "tactical": True}
    assert list(gov01["maturity"]) == ["0", "1", "2", "3", "4", "5"]
    assert gov01["profiles"] == ["Community Derived", "CORE Fundamentals"]
    assert list(gov01["possible_solutions"]) == [
        "micro_small",
        "small",
        "medium",
        "large",
        "enterprise",
    ]
    assert gov01["risks"] == ["R-AC-1", "R-GV-1"]
    assert gov01["threats"] == ["NT-1", "MT-1"]
    assert gov01["errata"] == "- reworded"
    assert "errata" not in controls[1]
    # Empty summary cells parse to real empty lists (the TDA-11.2 class).
    gov02 = controls[2]
    assert gov02["risks"] == [] and gov02["threats"] == []


def test_parse_crosswalks(clean_workbook):
    result = parse_workbook(clean_workbook, VERSION)
    xwalk = result.files["scf-crosswalks.json"]

    meta = xwalk["metadata"]
    assert meta["source"] == "Secure Controls Framework (SCF) v2099.1.1"
    assert meta["total_frameworks"] == 5  # orphan source matches no column
    assert meta["framework_display_names"]["general-nist-800-53-r5"] == "NIST SP 800-53 R5"
    # FDI collision gets the -2 suffix in sheet order.
    assert "general-dup-fw" in xwalk["scf_to_framework"]
    assert "general-dup-fw-2" in xwalk["scf_to_framework"]
    assert "usa-state-zz-orphan" not in xwalk["scf_to_framework"]

    fwd = xwalk["scf_to_framework"]["general-aicpa-tsc-2017"]
    assert fwd == {"GOV-01": ["CC1.1", "CC1.2"], "GOV-02": ["CC1.2"], "TST-01": ["CC9.9"]}
    rev = xwalk["framework_to_scf"]["general-aicpa-tsc-2017"]
    # Reverse map keys appear in first-appearance order; values in control order.
    assert list(rev) == ["CC1.1", "CC1.2", "CC9.9"]
    assert rev["CC1.2"] == ["GOV-01", "GOV-02"]


def test_parse_satellite_files(clean_workbook):
    result = parse_workbook(clean_workbook, VERSION)

    aos = result.files["scf-assessment-objectives.json"]
    assert aos["total"] == 3
    assert aos["assessment_objectives"][0] == {
        "scf_control_id": "GOV-01",
        "ao_id": "GOV-01_A01",
        "objective": "governance is governed.",
        "pptdf": "Process",
        "origin": "SCF Created",
        "assessment_rigor": "1",
        "scf_defined_parameters": "",
        "org_defined_parameters": "",
    }
    assert aos["assessment_objectives"][2]["assessment_rigor"] == "NIST 800-171"

    cc = result.files["scf-compensating-controls.json"]
    # The all-N/A row (GOV-01) is dropped.
    assert cc["total"] == 1
    rec = cc["compensating_controls"][0]
    assert rec["control_id"] == "GOV-01.1"
    assert rec["compensating_control_1"]["control_id"] == "GOV-02"
    assert rec["compensating_control_2"]["control_id"] == "MON-01"

    er = result.files["scf-evidence-requests.json"]
    assert er["total"] == 4
    assert er["evidence_requests"][1]["scf_controls"] == ["GOV-01", "GOV-02"]
    assert er["evidence_requests"][1]["cmmc_mapping"] == "AC.L2-3.1.1"
    assert er["evidence_requests"][0]["cmmc_mapping"] == ""

    pp = result.files["scf-privacy-principles.json"]
    assert pp["total"] == 2
    first, second = pp["privacy_principles"]
    assert first["principle_name"] == "Privacy by Design"
    assert first["crosswalks"] == {
        "AICPA TSC 2017:2022 (used for SOC 2)": ["CC1.1", "CC1.2"],
        "EU GDPR": ["Art 25"],
    }
    # Merged principle cells forward-fill.
    assert second["principle_name"] == "Privacy by Design"
    assert second["scf_control_id"] == "GOV-02"
    assert second["crosswalks"] == {"EU GDPR": ["Art 5"]}

    risks = result.files["scf-risks.json"]
    assert risks["total"] == 2
    assert risks["risks"][0] == {
        "risk_id": "R-AC-1",
        "grouping": "Access Control",
        "name": "No accountability",
        "description": "Cannot attribute.",
        "nist_csf_function": "Protect",
    }

    threats = result.files["scf-threats.json"]
    assert threats["total"] == 3
    # Grouping forward-fills across merged rows.
    assert threats["threats"][1] == {
        "threat_id": "NT-2",
        "grouping": "Natural Threat",
        "name": "Fire",
        "description": "Fire everywhere.",
    }


def test_parse_deterministic(clean_workbook):
    from scf_api import jsonio

    a = parse_workbook(clean_workbook, VERSION)
    b = parse_workbook(clean_workbook, VERSION)
    assert {n: jsonio.dumps(o) for n, o in a.files.items()} == {
        n: jsonio.dumps(o) for n, o in b.files.items()
    }


def test_renamed_tabs_and_headers_parse_identically(workbook_factory):
    """The 2026.2 rename shape yields byte-identical data to the old names.

    This is the regression that broke autopilot: a pure rename must change
    nothing about the emitted dataset.
    """
    from scf_api import jsonio

    old = parse_workbook(workbook_factory(), VERSION)
    renamed = parse_workbook(workbook_factory("renamed-tab", "renamed-cc-header"), VERSION)

    assert {n: jsonio.dumps(o) for n, o in old.files.items()} == {
        n: jsonio.dumps(o) for n, o in renamed.files.items()
    }
    # Built-in patterns handled it, so no operator alias was needed.
    assert renamed.findings == []


def test_operator_alias_surfaces_as_a_warn_finding(workbook_factory):
    """An overrides.json alias parses the workbook but flags the drift."""
    from scf_api.overrides import load_overrides
    from scf_api.validate import validate_result
    from scf_api.workbook import ResolutionAliases

    aliases = ResolutionAliases(sheets={"sources": [("reference register", "ref-register-2099")]})
    result = parse_workbook(workbook_factory("unknown-tab"), VERSION, aliases=aliases)

    (note,) = result.findings
    assert note["check"] == "resolution-alias"

    outcome = validate_result(result, load_overrides(None))
    alias_findings = [f for f in outcome.findings if f.check == "resolution-alias"]
    assert len(alias_findings) == 1
    assert alias_findings[0].severity == "warn"
    # A WARN must not block the gate.
    assert outcome.ok


# ---------------------------------------------------------------------------
# SCF 2026.3 layout: likelihood matrix, inline compensating controls
# ---------------------------------------------------------------------------


def test_2026_3_matrix_becomes_lists_and_likelihood_maps(workbook_factory):
    result = parse_workbook(workbook_factory("layout-2026-3"), VERSION)
    controls = {c["control_id"]: c for c in result.catalog["controls"]}

    gov01 = controls["GOV-01"]
    # Possible/Likely ids keep the pre-2026.3 meaning of the id lists, in column order.
    assert gov01["risks"] == ["R-AC-1", "R-GV-1"]
    assert gov01["threats"] == ["NT-1", "MT-1"]
    # The maps keep every rating, Unlikely included.
    assert gov01["risk_likelihood"] == {"R-AC-1": "Likely", "R-GV-1": "Possible"}
    assert gov01["threat_likelihood"] == {"NT-1": "Likely", "NT-2": "Unlikely", "MT-1": "Possible"}

    gov02 = controls["GOV-02"]
    assert gov02["risks"] == [] and gov02["threats"] == []
    assert set(gov02["risk_likelihood"].values()) == {"Unlikely"}


def test_2026_3_matches_legacy_layout_everywhere_but_the_likelihood_maps(workbook_factory):
    """One parser, two layouts: the same synthetic content yields the same data."""
    old = parse_workbook(workbook_factory(), VERSION)
    new = parse_workbook(workbook_factory("layout-2026-3"), VERSION)

    for name in old.files:
        if name == old.catalog_name:
            continue
        assert new.files[name] == old.files[name], name

    for o, n in zip(old.catalog["controls"], new.catalog["controls"], strict=True):
        n = {k: v for k, v in n.items() if k not in ("risk_likelihood", "threat_likelihood")}
        assert n == o


def test_2026_3_reads_compensating_controls_from_the_controls_tab(workbook_factory):
    result = parse_workbook(workbook_factory("layout-2026-3"), VERSION)
    cc = result.files["scf-compensating-controls.json"]
    assert cc["total"] == 1
    record = cc["compensating_controls"][0]
    assert record["control_id"] == "GOV-01.1"
    assert record["compensating_control_1"]["control_id"] == "GOV-02"
    assert record["compensating_control_2"]["control_id"] == "MON-01"


def test_2026_3_risk_name_survives_the_old_risk_id_column(workbook_factory):
    result = parse_workbook(workbook_factory("layout-2026-3"), VERSION)
    risks = result.files["scf-risks.json"]["risks"]
    assert risks[0]["risk_id"] == "R-AC-1"
    assert risks[0]["name"] == "No accountability"


def test_unknown_likelihood_rating_fails_the_gate(workbook_factory, tmp_path):
    from openpyxl import load_workbook

    from scf_api.overrides import load_overrides
    from scf_api.validate import validate_result

    path = workbook_factory("layout-2026-3")
    wb = load_workbook(path)
    ws = wb["SCF 2099.1"]
    col = next(c.column for c in ws[1] if c.value == "Risk R-AC-1")
    ws.cell(row=2, column=col, value="Probable")
    wb.save(path)

    outcome = validate_result(parse_workbook(path, VERSION), load_overrides(None))
    assert not outcome.ok
    assert any(f.check == "likelihood-value" and f.value == "Probable" for f in outcome.findings)
