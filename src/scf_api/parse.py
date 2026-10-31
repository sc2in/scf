"""Sheet extractors: SCF workbook -> the eight data/*.json structures.

Field order inside every record is part of the output contract (consumers
diff these files); dict insertion order below is deliberate. Output values
are raw cell text — normalization is for validation matching only.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl.utils import get_column_letter

from .normalize import norm_key, split_multivalue
from .sources import (
    MAPPING_LEFT_MARKER,
    MAPPING_RIGHT_MARKER,
    Reconciliation,
    SourceRow,
    assign_api_ids,
    matrix_columns,
    reconcile,
)
from .workbook import (
    ColumnResolveError,
    ResolutionAliases,
    SCFWorkbook,
    Sheet,
    cell_text,
)

SOURCE_URL = "https://securecontrolsframework.com"
LICENSE_TEXT = "CC BY-ND (Attribution, No Derivatives)"

# 2026.3 likelihood-matrix ratings. "Unlikely" means the risk/threat does not
# meaningfully apply, so it stays out of the `risks`/`threats` id lists (which
# keep their pre-2026.3 meaning: applicable ids) but is kept in the maps.
LIKELIHOODS = ("Unlikely", "Possible", "Likely")
APPLICABLE_LIKELIHOODS = frozenset({"Possible", "Likely"})

# 2026.3 added a `Legacy SCF #` column (previous-release ids of each control).
# `NONE` marks a new control; several ids (one per line) mark a merge.
LEGACY_MAP_FILE = "scf-legacy-map.json"
LEGACY_NONE = "none"


@dataclass
class Ref:
    """One cross-sheet reference, with workbook provenance for findings."""

    check: str  # validation check name, e.g. "control-erl"
    value: str  # raw referenced id text
    owner: str  # id of the referencing record
    sheet: str
    cell: str


@dataclass
class IdOccurrence:
    value: str
    sheet: str
    cell: str


@dataclass(frozen=True)
class PreviousRelease:
    """The release a workbook's `Legacy SCF #` ids refer to.

    `control_ids` is that release's full control-id set when known; the
    validation gate then requires the legacy map's old ids to equal it exactly.
    """

    version: str
    control_ids: frozenset[str] | None = None


@dataclass
class ParseResult:
    version: str
    files: dict[str, dict]  # data filename -> JSON-ready object
    refs: list[Ref] = field(default_factory=list)
    ids: dict[str, list[IdOccurrence]] = field(default_factory=dict)  # entity -> occurrences
    manifest: dict = field(default_factory=dict)  # sources reconciliation manifest
    workbook_sha256: str = ""
    findings: list[dict] = field(default_factory=list)  # sheet/col-resolve failures
    previous: PreviousRelease | None = None  # what legacy ids refer to

    @property
    def dashed(self) -> str:
        return self.version.replace(".", "-")

    @property
    def catalog_name(self) -> str:
        return f"scf-{self.dashed}.json"

    @property
    def catalog(self) -> dict:
        return self.files[self.catalog_name]


def _cell(sheet: Sheet, row_offset: int, col: int) -> str:
    """A1-style provenance for data row `row_offset` (0-based) in `sheet`."""
    return f"{get_column_letter(col + 1)}{sheet.header_row + 1 + row_offset}"


def _text(value: object) -> str:
    return cell_text(value) or ""


def parse_workbook(
    path: str | Path,
    version: str,
    aliases: ResolutionAliases | None = None,
    previous: PreviousRelease | None = None,
) -> ParseResult:
    path = Path(path)
    result = ParseResult(version=version, files={}, previous=previous)
    result.workbook_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()

    with SCFWorkbook(str(path), aliases=aliases) as wb:
        families = _parse_domains(wb, result)
        controls, xwalk = _parse_controls(wb, result)
        result.files[result.catalog_name] = {
            "framework": {
                "id": f"scf-{result.dashed}",
                "name": "SCF",
                "version": version,
            },
            "families": families,
            "controls": controls,
        }
        result.files["scf-crosswalks.json"] = xwalk
        if any("legacy_ids" in c for c in controls):
            deprecated = _parse_deprecated(wb, result)
            result.files[LEGACY_MAP_FILE] = _legacy_map(controls, deprecated, result)
        result.files["scf-assessment-objectives.json"] = _parse_aos(wb, result)
        result.files["scf-compensating-controls.json"] = _parse_compensating(wb, result)
        result.files["scf-evidence-requests.json"] = _parse_evidence(wb, result)
        result.files["scf-privacy-principles.json"] = _parse_privacy(wb, result)
        result.files["scf-risks.json"] = _parse_risks(wb, result)
        result.files["scf-threats.json"] = _parse_threats(wb, result)
        # Operator aliases that had to be used — surfaced as WARN findings.
        result.findings.extend(wb.resolution_notes)

    return result


# ---------------------------------------------------------------------------
# Domains
# ---------------------------------------------------------------------------


def _parse_domains(wb: SCFWorkbook, result: ParseResult) -> dict[str, str]:
    sheet = wb.sheet("domains")
    code_col = sheet.col("scf identifier")
    name_col = sheet.col("scf domain")
    families: dict[str, str] = {}
    occurrences = result.ids.setdefault("family", [])
    for i, row in enumerate(sheet.rows):
        code = cell_text(row[code_col]) if code_col < len(row) else None
        if not code:
            continue
        families[code.strip()] = _text(row[name_col])
        occurrences.append(IdOccurrence(code.strip(), sheet.title, _cell(sheet, i, code_col)))
    return families


# ---------------------------------------------------------------------------
# Controls + crosswalks (share the controls sheet scan)
# ---------------------------------------------------------------------------

_MATURITY_PATTERNS = [
    ("0", ("cmm level 0", "cmm 0")),
    ("1", ("cmm level 1", "cmm 1")),
    ("2", ("cmm level 2", "cmm 2")),
    ("3", ("cmm level 3", "cmm 3")),
    ("4", ("cmm level 4", "cmm 4")),
    ("5", ("cmm level 5", "cmm 5")),
]

_SOLUTION_KEYS = [
    ("micro_small", "micro-small"),
    ("small", "(10-49"),
    ("medium", "(50-249"),
    ("large", "(250-999"),
    ("enterprise", "> 1,000"),
]


def _profile_name(header: str) -> str:
    lines = [ln.strip() for ln in header.split("\n") if ln.strip()]
    if lines and lines[0] == "SCF":
        lines = lines[1:]
    return " ".join(lines)


def _parse_controls(wb: SCFWorkbook, result: ParseResult) -> tuple[list[dict], dict]:
    sheet = wb.sheet("controls")

    id_col = sheet.col("scf #")
    title_col = sheet.col("scf control", exclude=("question", "description"))
    desc_col = sheet.col("control description")
    cadence_col = sheet.col("conformity")
    erl_col = sheet.col("evidence request list")
    question_col = sheet.col("control question")
    weight_col = sheet.col("relative control weighting")
    pptdf_col = sheet.col("pptdf")
    nist_col = sheet.col("nist csf")
    scrm_cols = {
        "strategic": sheet.col("tier 1"),
        "operational": sheet.col("tier 2"),
        "tactical": sheet.col("tier 3"),
    }
    maturity_cols: dict[str, int] = {}
    for level, patterns in _MATURITY_PATTERNS:
        for pat in patterns:
            idx = sheet.find_col(pat)
            if idx is not None:
                maturity_cols[level] = idx
                break
        else:
            raise ColumnResolveError(f"no maturity column for level {level} in {sheet.title!r}")
    solution_cols = {key: sheet.col(pat) for key, pat in _SOLUTION_KEYS}

    # Risks/threats: a per-id likelihood matrix (2026.3) or two summary text
    # columns (<= 2026.2). The crosswalk slice ends where either begins.
    risk_matrix, threat_matrix = matrix_columns(sheet.header)
    if risk_matrix:
        risks_col = min(risk_matrix)
        threats_col = None
    else:
        risks_col = sheet.col(MAPPING_RIGHT_MARKER)
        threats_col = sheet.col("control threat summary")
    errata_col = sheet.find_col("errata")
    # 2026.3+: previous-release ids. Absent before, so optional.
    legacy_col = sheet.find_col("legacy scf #")

    # Crosswalk landmark slice: first AICPA column .. first risk column.
    left = next(
        (i for i, h in enumerate(sheet.header) if h and MAPPING_LEFT_MARKER in h),
        None,
    )
    if left is None:
        raise ColumnResolveError(f"no {MAPPING_LEFT_MARKER!r} landmark column in {sheet.title!r}")
    xwalk_cols = {
        i: sheet.header[i] for i in range(left, risks_col) if sheet.header[i] is not None
    }

    # Profile columns: between the last maturity column and the landmark,
    # headed by a leading "SCF" line.
    profile_start = max(maturity_cols.values()) + 1
    profile_cols = [
        (i, _profile_name(sheet.header[i]))
        for i in range(profile_start, left)
        if sheet.header[i] and sheet.header[i].split("\n")[0].strip() == "SCF"
    ]

    controls: list[dict] = []
    id_occurrences = result.ids.setdefault("control", [])

    def marked(value: object) -> bool:
        return value is not None and str(value).strip() != ""

    rows = [(i, row) for i, row in enumerate(sheet.rows) if cell_text(row[id_col])]

    for i, row in rows:
        cid = _text(row[id_col]).strip()
        id_occurrences.append(IdOccurrence(cid, sheet.title, _cell(sheet, i, id_col)))

        evidence = split_multivalue(row[erl_col])
        for v in evidence:
            result.refs.append(Ref("control-erl", v, cid, sheet.title, _cell(sheet, i, erl_col)))
        risk_likelihood: dict[str, str] = {}
        threat_likelihood: dict[str, str] = {}
        if risk_matrix:
            risks = _read_matrix(
                result, sheet, i, row, cid, risk_matrix, "control-risk", risk_likelihood
            )
            threats = _read_matrix(
                result, sheet, i, row, cid, threat_matrix, "control-threat", threat_likelihood
            )
        else:
            risks = split_multivalue(row[risks_col])
            for v in risks:
                result.refs.append(
                    Ref("control-risk", v, cid, sheet.title, _cell(sheet, i, risks_col))
                )
            threats = split_multivalue(row[threats_col])
            for v in threats:
                result.refs.append(
                    Ref("control-threat", v, cid, sheet.title, _cell(sheet, i, threats_col))
                )
        result.refs.append(
            Ref("family-domain", cid.split("-")[0], cid, sheet.title, _cell(sheet, i, id_col))
        )

        weight_raw = row[weight_col]
        weight = int(float(weight_raw)) if weight_raw is not None else 0

        control = {
            "control_id": cid,
            "title": _text(row[title_col]),
            "family": cid.split("-")[0],
            "description": _text(row[desc_col]),
            "scf_question": _text(row[question_col]),
            "relative_weight": weight,
            "conformity_cadence": _text(row[cadence_col]),
            "evidence_requests": evidence,
            "pptdf": _text(row[pptdf_col]),
            "nist_csf_function": _text(row[nist_col]),
            "scrm_focus": {key: marked(row[col]) for key, col in scrm_cols.items()},
            # Levels without criteria are omitted, not emitted empty.
            "maturity": {
                level: _text(row[col])
                for level, col in maturity_cols.items()
                if cell_text(row[col])
            },
            "profiles": [name for col, name in profile_cols if marked(row[col])],
            # Size tiers without guidance are omitted, not emitted empty.
            "possible_solutions": {
                key: _text(row[col]) for key, col in solution_cols.items() if cell_text(row[col])
            },
            "risks": risks,
            "threats": threats,
        }
        if risk_matrix:
            control["risk_likelihood"] = risk_likelihood
            control["threat_likelihood"] = threat_likelihood
        if errata_col is not None and cell_text(row[errata_col]):
            control["errata"] = _text(row[errata_col])
        if legacy_col is not None:
            control["legacy_ids"] = legacy_ids(row[legacy_col])
        controls.append(control)

    xwalk = _parse_crosswalks(wb, result, sheet, rows, id_col, xwalk_cols)
    return controls, xwalk


def legacy_ids(cell: object) -> list[str]:
    """Parse one `Legacy SCF #` cell: `NONE` -> [], else the newline-split ids."""
    return [v for v in split_multivalue(cell) if v.casefold() != LEGACY_NONE]


def _parse_deprecated(wb: SCFWorkbook, result: ParseResult) -> dict[str, str]:
    """READ THIS tab: deprecated old id -> the new id it merged into.

    The tab carries two side-by-side tables that both have a `Legacy SCF #`
    header; the deprecated one is the `Legacy SCF #` right of the
    `Deprecated SCF Control Name` header.
    """
    if not wb.has_sheet("readme"):
        result.findings.append(
            {
                "check": "legacy-deprecated",
                "severity": "warn",
                "message": "no READ THIS tab: no control is flagged deprecated in the legacy map",
                "owner": "",
                "value": "",
            }
        )
        return {}
    sheet = wb.catalog_sheet("readme", "deprecated scf control")
    name_col = sheet.col("deprecated scf control")
    old_col = next(
        (
            i
            for i, h in enumerate(sheet.header)
            if i > name_col and h and "legacy scf #" in norm_key(h)
        ),
        None,
    )
    if old_col is None:
        raise ColumnResolveError(
            f"no 'legacy scf #' column right of {sheet.header[name_col]!r} in {sheet.title!r}",
            sheet_title=sheet.title,
            patterns=("legacy scf #",),
            header=sheet.header,
        )
    new_col = sheet.col("merged into")

    deprecated: dict[str, str] = {}
    for i, row in enumerate(sheet.rows):
        old = cell_text(row[old_col]) if old_col < len(row) else None
        if not old or not old.strip():
            continue
        new = _text(row[new_col]).strip() if new_col < len(row) else ""
        old = old.strip()
        if old in deprecated and deprecated[old] != new:
            result.findings.append(
                {
                    "check": "legacy-deprecated",
                    "severity": "fail",
                    "message": (
                        f"deprecated {old!r} listed as merged into both "
                        f"{deprecated[old]!r} and {new!r}"
                    ),
                    "owner": old,
                    "value": new,
                    "sheet": sheet.title,
                    "cell": _cell(sheet, i, new_col),
                }
            )
            continue
        deprecated[old] = new
    return deprecated


def _legacy_map(controls: list[dict], deprecated: dict[str, str], result: ParseResult) -> dict:
    """Old (previous-release) id -> where it lives now.

    Built from `controls[].legacy_ids`; the validation gate proves it is a
    function onto real controls covering the previous release exactly.
    `merged`: the new control absorbed more than one old id. `deprecated`: the
    READ THIS tab lists the old id as a removed duplicate.
    """
    mapping: dict[str, dict] = {}
    for c in controls:
        olds = c.get("legacy_ids", [])
        for old in olds:
            if old in mapping:
                continue  # a split; the gate reports it (legacy-map-split)
            mapping[old] = {
                "new": c["control_id"],
                "merged": len(olds) > 1,
                "deprecated": old in deprecated,
            }
    for old, new in deprecated.items():
        entry = mapping.get(old)
        if entry is None or entry["new"] != new:
            where = "no control" if entry is None else repr(entry["new"])
            result.findings.append(
                {
                    "check": "legacy-deprecated",
                    "severity": "fail",
                    "message": (
                        f"READ THIS says deprecated {old!r} merged into {new!r}, "
                        f"but the Legacy SCF # column puts it on {where}"
                    ),
                    "owner": old,
                    "value": new,
                }
            )
    if result.previous is None:
        result.findings.append(
            {
                "check": "legacy-from-version",
                "severity": "warn",
                "message": (
                    "previous release unknown: the legacy map's from_version is null and "
                    "its old-id coverage is unchecked (parse into a data/ dir holding the "
                    "previous catalog, or pass --previous-catalog)"
                ),
                "owner": "",
                "value": "",
            }
        )
    return {
        "from_version": result.previous.version if result.previous else None,
        "to_version": result.version,
        "map": mapping,
    }


def _read_matrix(
    result: ParseResult,
    sheet: Sheet,
    i: int,
    row: tuple,
    cid: str,
    columns: dict[int, str],
    check: str,
    ratings: dict[str, str],
) -> list[str]:
    """Read one control's likelihood cells; return the applicable ids.

    Every id column becomes a Ref so the validation gate still proves it exists
    in the catalog. A rating outside LIKELIHOODS is a `likelihood-value` fail:
    that is the shape of the 2026.3 break (a matrix read as id text), so a new
    vocabulary must summon a human instead of silently dropping ids.
    """
    applicable: list[str] = []
    for col, rid in columns.items():
        raw = cell_text(row[col]) if col < len(row) else None
        if raw is None or not raw.strip():
            continue
        rating = next((r for r in LIKELIHOODS if r.casefold() == raw.strip().casefold()), None)
        result.refs.append(Ref(check, rid, cid, sheet.title, _cell(sheet, i, col)))
        if rating is None:
            result.findings.append(
                {
                    "check": "likelihood-value",
                    "severity": "fail",
                    "message": f"{cid!r} rates {rid!r} {raw!r}, expected one of {LIKELIHOODS}",
                    "owner": cid,
                    "value": raw,
                    "sheet": sheet.title,
                    "cell": _cell(sheet, i, col),
                }
            )
            continue
        ratings[rid] = rating
        if rating in APPLICABLE_LIKELIHOODS:
            applicable.append(rid)
    return applicable


def _js_object_order(d: dict[str, list[str]]) -> dict[str, list[str]]:
    """Reorder keys the way a JavaScript object iterates them.

    The published dataset was serialized from JS objects: canonical
    array-index keys (no leading zeros, < 2^32-1) come first in ascending
    numeric order, all other keys follow in insertion order. Framework refs
    are often bare numbers ("13", "7001"), so reverse crosswalk maps carry
    this ordering — consumers diff these files, so we reproduce it.
    """

    def is_index(k: str) -> bool:
        return k.isdigit() and str(int(k)) == k and int(k) < 2**32 - 1

    ordered = sorted((k for k in d if is_index(k)), key=int)
    ordered += [k for k in d if not is_index(k)]
    return {k: d[k] for k in ordered}


def _parse_crosswalks(
    wb: SCFWorkbook,
    result: ParseResult,
    controls_sheet: Sheet,
    control_rows: list[tuple[int, tuple]],
    id_col: int,
    xwalk_cols: dict[int, str],
) -> dict:
    sheet = wb.sheet("sources")
    header_col = sheet.col("scf column header")
    fdi_col = sheet.col("focal document identifier")
    geo_col = sheet.find_col("geography")
    source_col = sheet.find_col("source", exclude=("focal",))
    fdn_col = sheet.col("focal document name")

    source_rows: list[SourceRow] = []
    for i, row in enumerate(sheet.rows):
        fdi = cell_text(row[fdi_col]) if fdi_col < len(row) else None
        if not fdi:
            continue
        source_rows.append(
            SourceRow(
                row=sheet.header_row + 1 + i,
                geography=cell_text(row[geo_col]) if geo_col is not None else None,
                header=_text(row[header_col]),
                fdi=fdi.strip(),
                source=cell_text(row[source_col]) if source_col is not None else None,
                fdn=cell_text(row[fdn_col]),
            )
        )
    assign_api_ids(source_rows)

    recon: Reconciliation = reconcile(source_rows, xwalk_cols)
    result.manifest = recon.manifest()

    display_names: dict[str, str] = {}
    scf_to_framework: dict[str, dict] = {}
    framework_to_scf: dict[str, dict] = {}

    for match in recon.matches:
        if match.column_index is None:
            continue
        api_id = match.source.api_id
        display_names[api_id] = match.source.fdn or match.source.header
        forward: dict[str, list[str]] = {}
        reverse: dict[str, list[str]] = {}
        col = match.column_index
        for i, row in control_rows:
            cid = _text(row[id_col]).strip()
            refs = split_multivalue(row[col]) if col < len(row) else []
            if not refs:
                continue
            forward[cid] = refs
            for ref in refs:
                reverse.setdefault(ref, []).append(cid)
        scf_to_framework[api_id] = forward
        framework_to_scf[api_id] = _js_object_order(reverse)

    return {
        "metadata": {
            "source": f"Secure Controls Framework (SCF) v{result.version}",
            "source_url": SOURCE_URL,
            "license": LICENSE_TEXT,
            "total_frameworks": len(scf_to_framework),
            "framework_display_names": display_names,
        },
        "scf_to_framework": scf_to_framework,
        "framework_to_scf": framework_to_scf,
    }


# ---------------------------------------------------------------------------
# Assessment objectives
# ---------------------------------------------------------------------------


def _parse_aos(wb: SCFWorkbook, result: ParseResult) -> dict:
    sheet = wb.sheet("aos")
    control_col = sheet.col("scf #", exclude=("ao",))
    ao_col = sheet.col("scf ao #")
    objective_col = sheet.col("scf assessment objective", exclude=("origin",))
    pptdf_col = sheet.col("pptdf")
    origin_col = sheet.col("origin")
    rigor_col = sheet.col("rigor")
    sdp_col = sheet.col("scf defined parameters")
    # 2026.3 dropped the ODP column; "Notes" carries the ODP values (DoD ODP text).
    odp_col = sheet.col("organization defined parameters", alts=("notes",))

    records = []
    occurrences = result.ids.setdefault("ao", [])
    for i, row in enumerate(sheet.rows):
        ao_id = cell_text(row[ao_col]) if ao_col < len(row) else None
        if not ao_id:
            continue
        cid = _text(row[control_col]).strip()
        occurrences.append(IdOccurrence(ao_id.strip(), sheet.title, _cell(sheet, i, ao_col)))
        result.refs.append(
            Ref("ao-control", cid, ao_id.strip(), sheet.title, _cell(sheet, i, control_col))
        )
        records.append(
            {
                "scf_control_id": cid,
                "ao_id": ao_id.strip(),
                "objective": _text(row[objective_col]),
                "pptdf": _text(row[pptdf_col]),
                "origin": _text(row[origin_col]),
                "assessment_rigor": _text(row[rigor_col]),
                "scf_defined_parameters": _text(row[sdp_col]),
                "org_defined_parameters": _text(row[odp_col]),
            }
        )
    return {"total": len(records), "assessment_objectives": records}


# ---------------------------------------------------------------------------
# Compensating controls
# ---------------------------------------------------------------------------


def _parse_compensating(wb: SCFWorkbook, result: ParseResult) -> dict:
    # 2026.3 dropped the compensating tab: its columns now sit on the controls
    # tab, under the same headers.
    kind = "compensating" if wb.has_sheet("compensating") else "controls"
    sheet = wb.sheet(kind)
    # 2026.2 shortened "SCF Control #" to "SCF #". Neither pattern is a
    # substring of the other's header, so both releases resolve unambiguously.
    id_col = sheet.col("scf #", alts=("scf control #",))
    risk_col = sheet.col("risk if")
    cc1_col = sheet.col("possible compensating control #1")
    cc2_col = sheet.col("possible compensating control #2")

    def cc_block(row: tuple, base: int, cid: str, i: int) -> dict | None:
        ref = cell_text(row[base]) if base < len(row) else None
        if not ref or ref.strip() in ("N/A", "NA"):
            return None
        ref = ref.strip()
        result.refs.append(Ref("cc-control", ref, cid, sheet.title, _cell(sheet, i, base)))
        return {
            "control_id": ref,
            "name": _text(row[base + 1]),
            "description": _text(row[base + 2]),
            "justification": _text(row[base + 3]),
        }

    records = []
    for i, row in enumerate(sheet.rows):
        cid = cell_text(row[id_col]) if id_col < len(row) else None
        if not cid:
            continue
        cid = cid.strip()
        cc1 = cc_block(row, cc1_col, cid, i)
        cc2 = cc_block(row, cc2_col, cid, i)
        if cc1 is None and cc2 is None:
            continue
        result.refs.append(Ref("cc-control", cid, cid, sheet.title, _cell(sheet, i, id_col)))
        record = {
            "control_id": cid,
            "risk_if_not_implemented": _text(row[risk_col]),
            "compensating_control_1": cc1,
            "compensating_control_2": cc2,
        }
        records.append(record)
    return {"total": len(records), "compensating_controls": records}


# ---------------------------------------------------------------------------
# Evidence requests
# ---------------------------------------------------------------------------


def _parse_evidence(wb: SCFWorkbook, result: ParseResult) -> dict:
    sheet = wb.sheet("evidence")
    erl_col = sheet.col("erl #")
    area_col = sheet.col("area of focus")
    # 2026.3 renamed "Documentation Artifact" to "ERL Artifact".
    artifact_col = sheet.col("erl artifact", alts=("documentation artifact",))
    desc_col = sheet.col("artifact description")
    mappings_col = sheet.col("scf control mappings")
    cmmc_col = sheet.find_col("cmmc")

    records = []
    occurrences = result.ids.setdefault("erl", [])
    for i, row in enumerate(sheet.rows):
        erl_id = cell_text(row[erl_col]) if erl_col < len(row) else None
        if not erl_id:
            continue
        erl_id = erl_id.strip()
        occurrences.append(IdOccurrence(erl_id, sheet.title, _cell(sheet, i, erl_col)))
        scf_controls = split_multivalue(row[mappings_col])
        for v in scf_controls:
            result.refs.append(
                Ref("erl-control", v, erl_id, sheet.title, _cell(sheet, i, mappings_col))
            )
        records.append(
            {
                "erl_id": erl_id,
                "area": _text(row[area_col]),
                "artifact_name": _text(row[artifact_col]),
                "artifact_description": _text(row[desc_col]),
                "scf_controls": scf_controls,
                "cmmc_mapping": _text(row[cmmc_col]) if cmmc_col is not None else "",
            }
        )
    return {"total": len(records), "evidence_requests": records}


# ---------------------------------------------------------------------------
# Privacy principles
# ---------------------------------------------------------------------------


def _parse_privacy(wb: SCFWorkbook, result: ParseResult) -> dict:
    sheet = wb.sheet("privacy")
    name_col = sheet.col("principle name")
    desc_col = sheet.col("scf-dpmp")
    control_name_col = sheet.col("scf control", exclude=("description", "#"))
    control_id_col = sheet.col("scf #")
    first_xwalk = control_id_col + 2  # skip the SCF control description column

    xwalk_headers = {
        i: sheet.header[i].replace("\n", " ")
        for i in range(first_xwalk, len(sheet.header))
        if sheet.header[i]
    }

    records = []
    name = desc = ""
    for i, row in enumerate(sheet.rows):
        if cell_text(row[name_col]):
            name = _text(row[name_col])
        if cell_text(row[desc_col]):
            desc = _text(row[desc_col])
        cid = cell_text(row[control_id_col]) if control_id_col < len(row) else None
        if not cid:
            continue
        cid = cid.strip()
        result.refs.append(
            Ref("privacy-control", cid, name, sheet.title, _cell(sheet, i, control_id_col))
        )
        crosswalks = {}
        for col, header in xwalk_headers.items():
            refs = split_multivalue(row[col]) if col < len(row) else []
            if refs:
                crosswalks[header] = refs
        records.append(
            {
                "principle_name": name,
                "description": desc,
                "scf_control": _text(row[control_name_col]),
                "scf_control_id": cid,
                "crosswalks": crosswalks,
            }
        )
    return {"total": len(records), "privacy_principles": records}


# ---------------------------------------------------------------------------
# Risk & threat catalogs
# ---------------------------------------------------------------------------


def _parse_risks(wb: SCFWorkbook, result: ParseResult) -> dict:
    sheet = wb.catalog_sheet("risks", "Risk #")
    id_col = sheet.col("risk #")
    grouping_col = sheet.col("grouping")
    desc_col = sheet.col("description")
    func_col = sheet.col("nist csf")
    # 2026.3 inserted "Old Risk #" after "Risk #", so the name is no longer the
    # next column; it is the "Risk*" header.
    name_col = sheet.find_col("risk*")
    if name_col is None:
        name_col = id_col + 1

    records = []
    occurrences = result.ids.setdefault("risk", [])
    grouping = ""
    for i, row in enumerate(sheet.rows):
        if grouping_col < len(row) and cell_text(row[grouping_col]):
            grouping = _text(row[grouping_col])
        rid = cell_text(row[id_col]) if id_col < len(row) else None
        if not rid:
            continue
        rid = rid.strip()
        occurrences.append(IdOccurrence(rid, sheet.title, _cell(sheet, i, id_col)))
        records.append(
            {
                "risk_id": rid,
                "grouping": grouping,
                "name": _text(row[name_col]),
                "description": _text(row[desc_col]),
                "nist_csf_function": _text(row[func_col]),
            }
        )
    return {"total": len(records), "risks": records}


def _parse_threats(wb: SCFWorkbook, result: ParseResult) -> dict:
    sheet = wb.catalog_sheet("threats", "Threat #")
    id_col = sheet.col("threat #")
    grouping_col = sheet.col("grouping")
    desc_col = sheet.col("description")
    name_col = id_col + 1

    records = []
    occurrences = result.ids.setdefault("threat", [])
    grouping = ""
    for i, row in enumerate(sheet.rows):
        if grouping_col < len(row) and cell_text(row[grouping_col]):
            grouping = _text(row[grouping_col])
        tid = cell_text(row[id_col]) if id_col < len(row) else None
        if not tid:
            continue
        tid = tid.strip()
        occurrences.append(IdOccurrence(tid, sheet.title, _cell(sheet, i, id_col)))
        records.append(
            {
                "threat_id": tid,
                "grouping": grouping,
                "name": _text(row[name_col]),
                "description": _text(row[desc_col]),
            }
        )
    return {"total": len(records), "threats": records}
