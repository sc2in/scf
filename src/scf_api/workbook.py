"""Workbook access: pattern-based sheet resolution and cell flattening.

Sheets are found by normalized substring pattern, never by index — sheet
order and exact names change between SCF releases (2026.2 renamed
"Authoritative Sources" to "Focal Documents"). Each logical kind therefore
carries an ORDERED tuple of accepted patterns: the current name first,
historical names after it, so one parser reads every release we support.

Patterns are matched with `norm_key`, the same normalization column headers
get, so a tab that gains an NBSP, an en-dash or a double space still
resolves instead of hard-failing the autopilot.

When neither the built-in patterns nor an operator-supplied alias resolves a
sheet, that is a `SheetResolveError` and the validation gate summons a human.

Cells are read with openpyxl in read_only+data_only mode, so formula cells
yield their cached values (the TEXTJOIN risk/threat summaries) and rich text
arrives as plain strings.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl import load_workbook

from .normalize import norm_key


class SheetResolveError(KeyError):
    """A required sheet could not be found by pattern."""


class ColumnResolveError(KeyError):
    """A required column could not be found in a sheet header."""


# kind -> (accepted patterns, newest first), exclude-patterns.
# All matched with norm_key (NFKC, dash-folded, casefolded, space-collapsed).
SHEET_PATTERNS: dict[str, tuple[tuple[str, ...], tuple[str, ...]]] = {
    "domains": (("domain",), ()),
    # 2026.2 renamed "Authoritative Sources" -> "Focal Documents".
    "sources": (("focal document", "authoritative sources"), ()),
    "controls": (
        ("scf",),
        (
            "domain",
            "authoritative",
            "focal",
            "privacy",
            "objective",
            "list",
            "compensating",
            "catalog",
        ),
    ),
    # Dropped in 2026.3 (folded into the controls tab); still present before.
    "compensating": (("compensating",), ()),
    "evidence": (("evidence request",), ()),
    "aos": (("assessment objectives",), ("evidence",)),
    "privacy": (("privacy",), ("evidence",)),
    "risks": (("risk catalog",), ()),
    "threats": (("threat catalog",), ()),
}


@dataclass
class ResolutionAliases:
    """Operator-supplied fallback patterns from overrides/overrides.json.

    Tried only after every built-in pattern has missed, so an alias can never
    steal a sheet from the shipped manifest. A hit is recorded as a WARN
    finding: the dataset still builds unattended, but the drift stays visible
    so the alias can be folded into SHEET_PATTERNS in a later release.
    """

    sheets: dict[str, list[tuple[str, str]]] = field(default_factory=dict)
    columns: dict[str, list[tuple[str, str, str]]] = field(default_factory=dict)

    def sheet_aliases(self, kind: str) -> list[tuple[str, str]]:
        """[(pattern, override-id)] for a logical sheet kind."""
        return self.sheets.get(kind, [])

    def column_aliases(self, kind: str, pattern: str) -> list[tuple[str, str]]:
        """[(alias, override-id)] for a (sheet kind, primary pattern) pair."""
        want = norm_key(pattern)
        return [
            (alias, oid) for pat, alias, oid in self.columns.get(kind, []) if norm_key(pat) == want
        ]


def cell_text(value: object) -> str | None:
    """Flatten a cell value to text. Numbers keep their shortest repr."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


@dataclass
class Sheet:
    """A materialized sheet: title, header row, and data rows (raw values)."""

    title: str
    header: list[str | None]
    rows: list[tuple]
    header_row: int  # 1-based row number of the header (provenance)
    kind: str = ""  # logical kind it was resolved as (alias lookup key)
    aliases: "ResolutionAliases | None" = None
    notes: list[dict] | None = None  # shared sink for resolution WARN findings

    def col(
        self,
        pattern: str,
        exclude: tuple[str, ...] = (),
        alts: tuple[str, ...] = (),
    ) -> int:
        """Resolve a column, trying `pattern` then each of `alts` in order."""
        idx = self.find_col(pattern, exclude, alts)
        if idx is None:
            tried = ", ".join(repr(p) for p in (pattern, *alts))
            raise ColumnResolveError(
                f"no column matching {tried} in sheet {self.title!r} "
                f"(headers: {[h for h in self.header if h][:12]}...)"
            )
        return idx

    def find_col(
        self,
        pattern: str,
        exclude: tuple[str, ...] = (),
        alts: tuple[str, ...] = (),
    ) -> int | None:
        """First matching column index, or None.

        Built-in patterns (`pattern`, then `alts`) resolve silently — they are
        the shipped, tested manifest. Operator aliases from overrides.json are
        tried last and record a WARN so the drift is visible in the report.
        """
        for pat in (pattern, *alts):
            idx = self._scan(pat, exclude)
            if idx is not None:
                return idx
        if self.aliases is not None:
            for alias, oid in self.aliases.column_aliases(self.kind, pattern):
                idx = self._scan(alias, exclude)
                if idx is not None:
                    self._note(
                        "column",
                        f"column {pattern!r} in sheet {self.title!r} resolved via "
                        f"operator alias {alias!r} (override {oid!r}) — fold it into "
                        f"the parser's built-in patterns",
                        oid,
                    )
                    return idx
        return None

    def _scan(self, pattern: str, exclude: tuple[str, ...]) -> int | None:
        want = norm_key(pattern)
        for i, h in enumerate(self.header):
            if h is None:
                continue
            key = norm_key(str(h))
            if want in key and not any(norm_key(x) in key for x in exclude):
                return i
        return None

    def _note(self, scope: str, message: str, oid: str) -> None:
        if self.notes is None:
            return
        self.notes.append(
            {
                "check": "resolution-alias",
                "severity": "warn",
                "message": message,
                "owner": f"{scope}:{self.kind}",
                "value": oid,
                "sheet": self.title,
            }
        )


class SCFWorkbook:
    """The SCF master workbook, resolved sheet by pattern."""

    def __init__(self, path: str, aliases: ResolutionAliases | None = None):
        self.path = path
        self._wb = load_workbook(path, read_only=True, data_only=True)
        self.sheet_names: list[str] = list(self._wb.sheetnames)
        self.aliases = aliases
        # Resolution WARNs (operator aliases used); drained into ParseResult.
        self.resolution_notes: list[dict] = []

    def close(self) -> None:
        self._wb.close()

    def __enter__(self) -> "SCFWorkbook":
        return self

    def __exit__(self, *exc: object) -> None:
        self.close()

    def resolve_sheet_name(self, kind: str) -> str:
        """Resolve a logical kind to a tab name.

        Built-in patterns are tried in order (current name first, historical
        names after) and resolve silently. Operator aliases from
        overrides.json are tried only after all of them miss, and record a
        WARN so a rename patched in production stays visible.
        """
        patterns, excludes = SHEET_PATTERNS[kind]
        for pattern in patterns:
            name = self._scan_sheets(pattern, excludes)
            if name is not None:
                return name

        if self.aliases is not None:
            for alias, oid in self.aliases.sheet_aliases(kind):
                name = self._scan_sheets(alias, excludes)
                if name is not None:
                    self.resolution_notes.append(
                        {
                            "check": "resolution-alias",
                            "severity": "warn",
                            "message": (
                                f"sheet {kind!r} resolved to {name!r} via operator alias "
                                f"{alias!r} (override {oid!r}) — fold it into "
                                f"SHEET_PATTERNS in the parser"
                            ),
                            "owner": f"sheet:{kind}",
                            "value": oid,
                            "sheet": name,
                        }
                    )
                    return name

        tried = ", ".join(repr(p) for p in patterns)
        raise SheetResolveError(
            f"no sheet matching {tried} (exclude={excludes!r}) in {self.sheet_names}"
        )

    def has_sheet(self, kind: str) -> bool:
        """True when a built-in pattern resolves `kind` (no aliases, never raises).

        Used for sheets upstream has dropped: 2026.3 folded the compensating
        controls tab into the controls tab, so its absence is a layout, not an error.
        """
        patterns, excludes = SHEET_PATTERNS[kind]
        return any(self._scan_sheets(p, excludes) is not None for p in patterns)

    def _scan_sheets(self, pattern: str, excludes: tuple[str, ...]) -> str | None:
        want = norm_key(pattern)
        for name in self.sheet_names:
            key = norm_key(name)
            if want in key and not any(norm_key(x) in key for x in excludes):
                return name
        return None

    def sheet(self, kind: str, header_row: int = 1) -> Sheet:
        """Materialize a sheet. header_row is 1-based; rows are what follows it."""
        name = self.resolve_sheet_name(kind)
        ws = self._wb[name]
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < header_row:
            raise ColumnResolveError(f"sheet {name!r} has no header row {header_row}")
        header = [cell_text(c) for c in all_rows[header_row - 1]]
        return Sheet(
            title=name,
            header=header,
            rows=all_rows[header_row:],
            header_row=header_row,
            kind=kind,
            aliases=self.aliases,
            notes=self.resolution_notes,
        )

    def catalog_sheet(self, kind: str, id_header: str) -> Sheet:
        """Resolve a catalog-style sheet (Risk/Threat) whose header row floats.

        These sheets carry several rows of prose before the real header; find
        the first row containing id_header and treat it as the header row.
        """
        name = self.resolve_sheet_name(kind)
        ws = self._wb[name]
        all_rows = list(ws.iter_rows(values_only=True))
        want = norm_key(id_header)
        for i, row in enumerate(all_rows):
            if any(c is not None and want in norm_key(str(c)) for c in row):
                return Sheet(
                    title=name,
                    header=[cell_text(c) for c in row],
                    rows=all_rows[i + 1 :],
                    header_row=i + 1,
                    kind=kind,
                    aliases=self.aliases,
                    notes=self.resolution_notes,
                )
        raise ColumnResolveError(f"no header row containing {id_header!r} in sheet {name!r}")
